import io
from decimal import Decimal

import openpyxl
from fastapi.testclient import TestClient

from app.services.skoring import hitung_skor_kebijakan


def test_rekap_api_endpoint(client: TestClient) -> None:
    client.post(
        "/api/v1/auth/masuk",
        json={"username": "koordinator", "password": "Password123!"},
    )
    response = client.get("/api/v1/rekap")
    assert response.status_code == 200
    data = response.json()
    assert "periode" in data
    assert "kebijakan_list" in data
    assert len(data["matriks_kesiapan"]) == 11
    assert data["periode"]["tahun"] == 2026


def test_skoring_parsial_fallback() -> None:
    dimensi_cfg = [
        {
            "kode": "A",
            "nama": "Perencanaan",
            "bobot": "25",
            "pertanyaan": [
                {"kode": "A1", "bobot_relatif": "1"},
                {"kode": "A2", "bobot_relatif": "1"},
            ],
        }
    ]
    # Hanya A1 dijawab, A2 belum
    jawaban = {"A1": Decimal("100")}
    hasil = hitung_skor_kebijakan(
        melibatkan_jfak=True,
        bobot_profil=Decimal("10"),
        skor_profil_ya=Decimal("100"),
        skor_profil_tidak=None,
        dimensi_config=dimensi_cfg,
        jawaban_skor=jawaban,
        izinkan_parsial=True,
    )
    # A1 = 100, A2 = 0 -> rata-rata A = 50 -> kontribusi A = 25% * 50 = 12.5
    # Profil = 10% * 100 = 10 -> Total = 22.5
    assert hasil.dimensi["A"].nilai_murni == Decimal("50")
    assert hasil.dimensi["A"].kontribusi_total == Decimal("12.5")
    assert hasil.nilai_total == Decimal("22.5")


def test_ekspor_excel_rekap(client: TestClient) -> None:
    client.post(
        "/api/v1/auth/masuk",
        json={"username": "koordinator", "password": "Password123!"},
    )
    response = client.get("/api/v1/rekap/ekspor-excel")
    assert response.status_code == 200
    assert (
        response.headers["content-type"]
        == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert "attachment; filename=" in response.headers["content-disposition"]

    wb = openpyxl.load_workbook(io.BytesIO(response.content))
    assert wb.sheetnames == ["Tabel Rekap"]

    ws = wb["Tabel Rekap"]
    assert ws["C4"].value == "A. Perencanaan Kebijakan (20)"
    assert ws["D4"].value == "B. Implementasi Kebijakan (25)"
    assert ws["E4"].value == "C. Evaluasi & Keberlanjutan Kebijakan (30)"
    assert ws["F4"].value == "D. Tranparansi & Partisipasi Publik (15)"

    # Periksa seluruh baris data kebijakan yang diekspor
    rows_data = []
    for r in range(5, ws.max_row + 1):
        if ws.cell(row=r, column=1).value is not None:
            rows_data.append(
                {
                    "no": ws.cell(row=r, column=1).value,
                    "nama": ws.cell(row=r, column=2).value,
                    "a": ws.cell(row=r, column=3).value,
                    "b": ws.cell(row=r, column=4).value,
                    "c": ws.cell(row=r, column=5).value,
                    "d": ws.cell(row=r, column=6).value,
                    "total": ws.cell(row=r, column=7).value,
                }
            )
    print("DATA ROWS IN EXCEL:", rows_data)
    assert len(rows_data) >= 3
    # Pastikan setidaknya satu kebijakan memiliki skor dimensi yang valid dan lebih dari 0
    pse_row = next(r for r in rows_data if "PSE" in str(r["nama"]))
    assert float(pse_row["a"]) == 14.0
    assert float(pse_row["b"]) == 20.94
    assert float(pse_row["c"]) in (25.12, 25.13)
    assert float(pse_row["d"]) in (13.12, 13.13)
    assert float(pse_row["total"]) == 83.19
