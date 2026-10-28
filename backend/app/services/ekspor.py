import io
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, Side

from app.schemas.rekap import RekapPeriodeBaca


def buat_ekspor_excel_rekap(rekap: RekapPeriodeBaca) -> io.BytesIO:
    """
    Mengisi data rekapitulasi asesmen IKK ke dalam template Excel resmi:
    - Sheet 'Tabel Rekap': Memuat No, Nama Kebijakan, Kontribusi Dimensi A-D, Nilai Total,
      Kualifikasi, dan Rekomendasi sesuai format resmi LAN RI.
    - Sheet lain dihapus sehingga berkas hasil ekspor hanya berisi 'Tabel Rekap'.
    """
    # Cari berkas template di docs/referensi atau fallback di app/seeds
    base_dir = Path(__file__).resolve().parent.parent.parent.parent
    seeds_dir = Path(__file__).resolve().parent.parent / "seeds"
    kandidat_paths = [
        base_dir
        / "docs"
        / "referensi"
        / "Copy of Template Form Self Assessment Instrumen IKK 2026.xlsx",
        base_dir / "docs" / "referensi" / "template_form_self_assessment_ikk_2026.xlsx",
        seeds_dir / "template_form_self_assessment_ikk_2026.xlsx",
    ]
    template_path = next((p for p in kandidat_paths if p.exists()), None)

    if template_path is not None:
        wb = openpyxl.load_workbook(template_path)
    else:
        wb = openpyxl.Workbook()
        ws_init = wb.active
        if ws_init is not None:
            ws_init.title = "Tabel Rekap"
            ws_init["A4"] = "No"
            ws_init["B4"] = "Nama Kebijakan"
            ws_init["C4"] = "A. Perencanaan Kebijakan (20)"
            ws_init["D4"] = "B. Implementasi Kebijakan (25)"
            ws_init["E4"] = "C. Evaluasi & Keberlanjutan Kebijakan (30)"
            ws_init["F4"] = "D. Tranparansi & Partisipasi Publik (15)"
            ws_init["G4"] = "Nilai Total"
            ws_init["H4"] = "Kualifikasi"
            ws_init["I4"] = "Rekomendasi"

    # 1. Update Sheet 'Tabel Rekap'
    if "Tabel Rekap" in wb.sheetnames:
        ws_rekap = wb["Tabel Rekap"]

        # Perbarui bobot di header Row 4
        ws_rekap["C4"] = "A. Perencanaan Kebijakan (20)"
        ws_rekap["D4"] = "B. Implementasi Kebijakan (25)"
        ws_rekap["E4"] = "C. Evaluasi & Keberlanjutan Kebijakan (30)"
        ws_rekap["F4"] = "D. Tranparansi & Partisipasi Publik (15)"

        # Bersihkan baris data template lama mulai dari baris 5
        if ws_rekap.max_row >= 5:
            ws_rekap.delete_rows(5, ws_rekap.max_row - 4)

        # Style baku
        thin_side = Side(style="thin", color="000000")
        thin_border = Border(
            left=thin_side,
            right=thin_side,
            top=thin_side,
            bottom=thin_side,
        )
        font_regular = Font(name="Calibri", size=11, bold=False)
        font_bold = Font(name="Calibri", size=11, bold=True)

        align_center = Alignment(horizontal="center", vertical="center")
        align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)
        align_left_top = Alignment(horizontal="left", vertical="top", wrap_text=True)

        # Urutkan: Kebijakan terpilih untuk SI IKK di atas, disusul kandidat lainnya
        sorted_kebijakan = sorted(
            rekap.kebijakan_list,
            key=lambda k: (not k.dipilih_untuk_si_ikk, k.nama.lower()),
        )

        current_row = 5
        for idx, k in enumerate(sorted_kebijakan, start=1):
            # No
            c_no = ws_rekap.cell(row=current_row, column=1, value=idx)
            c_no.font = font_regular
            c_no.alignment = align_center
            c_no.border = thin_border

            # Nama Kebijakan
            nama_label = k.nama
            if k.dipilih_untuk_si_ikk:
                nama_label += " (Dipilih untuk SI IKK)"
            c_nama = ws_rekap.cell(row=current_row, column=2, value=nama_label)
            c_nama.font = font_regular
            c_nama.alignment = align_left
            c_nama.border = thin_border

            # Dimensi A
            dim_a = k.dimensi.get("a") or k.dimensi.get("A")
            val_a = float(dim_a.kontribusi_total) if dim_a else 0.0
            c_a = ws_rekap.cell(row=current_row, column=3, value=round(val_a, 2))
            c_a.font = font_regular
            c_a.alignment = align_center
            c_a.number_format = "#,##0.00"
            c_a.border = thin_border

            # Dimensi B
            dim_b = k.dimensi.get("b") or k.dimensi.get("B")
            val_b = float(dim_b.kontribusi_total) if dim_b else 0.0
            c_b = ws_rekap.cell(row=current_row, column=4, value=round(val_b, 2))
            c_b.font = font_regular
            c_b.alignment = align_center
            c_b.number_format = "#,##0.00"
            c_b.border = thin_border

            # Dimensi C
            dim_c = k.dimensi.get("c") or k.dimensi.get("C")
            val_c = float(dim_c.kontribusi_total) if dim_c else 0.0
            c_c = ws_rekap.cell(row=current_row, column=5, value=round(val_c, 2))
            c_c.font = font_regular
            c_c.alignment = align_center
            c_c.number_format = "#,##0.00"
            c_c.border = thin_border

            # Dimensi D
            dim_d = k.dimensi.get("d") or k.dimensi.get("D")
            val_d = float(dim_d.kontribusi_total) if dim_d else 0.0
            c_d = ws_rekap.cell(row=current_row, column=6, value=round(val_d, 2))
            c_d.font = font_regular
            c_d.alignment = align_center
            c_d.number_format = "#,##0.00"
            c_d.border = thin_border

            # Nilai Total
            val_total = (
                float(k.nilai_total_tampilan)
                if k.nilai_total_tampilan
                else float(k.nilai_total or 0.0)
            )
            c_total = ws_rekap.cell(row=current_row, column=7, value=round(val_total, 2))
            c_total.font = font_bold
            c_total.alignment = align_center
            c_total.number_format = "#,##0.00"
            c_total.border = thin_border

            # Kualifikasi
            kual_str = (k.kualifikasi or "Kurang").upper()
            c_kual = ws_rekap.cell(row=current_row, column=8, value=kual_str)
            c_kual.font = font_bold
            c_kual.alignment = align_center
            c_kual.border = thin_border

            # Rekomendasi
            c_rek = ws_rekap.cell(row=current_row, column=9, value=k.rekomendasi or "")
            c_rek.font = font_regular
            c_rek.alignment = align_left_top
            c_rek.border = thin_border

            ws_rekap.row_dimensions[current_row].height = 28
            current_row += 1

        # Atur lebar kolom agar proporsional dan nyaman dibaca
        ws_rekap.column_dimensions["A"].width = 6
        ws_rekap.column_dimensions["B"].width = 38
        ws_rekap.column_dimensions["C"].width = 16
        ws_rekap.column_dimensions["D"].width = 16
        ws_rekap.column_dimensions["E"].width = 16
        ws_rekap.column_dimensions["F"].width = 16
        ws_rekap.column_dimensions["G"].width = 14
        ws_rekap.column_dimensions["H"].width = 16
        ws_rekap.column_dimensions["I"].width = 45

    # Hapus semua sheet lain selain 'Tabel Rekap' agar berkas ekspor hanya berisi tabel rekap
    for sheetname in list(wb.sheetnames):
        if sheetname != "Tabel Rekap":
            wb.remove(wb[sheetname])

    wb.active = wb["Tabel Rekap"]

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output
